"""Export BaiYun_Agent audit log (jsonl) to Excel.

Default: concise version - only user question + AI final reply per round.
--detailed: also insert one row per tool_call with its visible_text.

Usage:
  python export_excel.py -i audit.jsonl [-u 用户名] -o out.xlsx [--detailed] [--img-dir ./images] [--desc "file.jpg:描述"] 

Images in user_input.image_urls (data:image base64) are saved to --img-dir
and embedded into the "图片预览" column.
"""
from __future__ import annotations

import argparse
import base64
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageOps

HEADERS = ["用户名称", "角色", "时间戳", "对话内容", "备注", "图片描述", "图片文件", "图片预览"]
WIDTHS = {"A": 14, "B": 12, "C": 20, "D": 80, "E": 22, "F": 58, "G": 38, "H": 92}
USER_FILL = PatternFill("solid", fgColor="E8F5E9")
AI_FILL = PatternFill("solid", fgColor="FFF3E0")
IMG_FILL = PatternFill("solid", fgColor="F3E5F5")
TOOL_FILL = PatternFill("solid", fgColor="E3F2FD")
HDR_FILL = PatternFill("solid", fgColor="D9EAF7")


def normalize_ts(value: str) -> str:
    text = (value or "").strip()
    if "T" in text:
        try:
            return datetime.fromisoformat(text).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            return text.replace("T", " ")[:19]
    return text[:19]


def strip_prefix(text: str, prefix: str = "conversation_id:") -> str:
    if text.lstrip().startswith(prefix):
        return text.split("\n", 1)[-1]
    return text


def build_preview(items: list[dict], output_path: Path) -> tuple[int, int]:
    limit = (330, 240) if len(items) > 1 else (460, 300)
    thumbs: list[Image.Image] = []
    for item in items:
        img = Image.open(item["path"])
        img = ImageOps.exif_transpose(img).convert("RGB")
        img.thumbnail(limit, Image.Resampling.LANCZOS)
        label = Image.new("RGB", (img.width, img.height + 22), "white")
        d = ImageDraw.Draw(label)
        d.rectangle((0, 0, label.width - 1, label.height - 1), outline=(190, 190, 190))
        d.text((6, 5), item["name"], fill=(70, 70, 70))
        label.paste(img, (0, 22))
        thumbs.append(label)
    gap = 12
    w = sum(t.width for t in thumbs) + gap * (len(thumbs) - 1)
    h = max(t.height for t in thumbs)
    canvas = Image.new("RGB", (w, h), "white")
    x = 0
    for t in thumbs:
        canvas.paste(t, (x, 0))
        x += t.width + gap
    canvas.save(output_path, quality=90)
    return canvas.size


def extract_images(user_input: dict, img_dir: Path, round_idx: int,
                   desc_map: dict) -> list[dict]:
    items = []
    for idx, url in enumerate(user_input.get("image_urls") or []):
        if not (isinstance(url, str) and url.startswith("data:image")):
            continue
        header, b64 = url.split(",", 1)
        ext = "png" if "png" in header else "jpg"
        name = f"round{round_idx}_{idx}.{ext}"
        path = img_dir / name
        path.write_bytes(base64.b64decode(b64))
        items.append({"name": name, "path": path, "desc": desc_map.get(name, "")})
    return items


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("-i", "--input", required=True, help="audit jsonl file")
    ap.add_argument("-o", "--output", required=True, help="output xlsx path")
    ap.add_argument("-u", "--user", default=None, help="filter by user_name")
    ap.add_argument("--detailed", action="store_true", help="include tool_call rows")
    ap.add_argument("--img-dir", default="./images", help="image output dir")
    ap.add_argument("--desc", action="append", default=[], help="file.jpg:描述")
    args = ap.parse_args()

    desc_map = {}
    for d in args.desc:
        if ":" in d:
            k, v = d.split(":", 1)
            desc_map[k] = v

    img_dir = Path(args.img_dir)
    img_dir.mkdir(parents=True, exist_ok=True)
    preview_dir = img_dir / "_previews"
    preview_dir.mkdir(exist_ok=True)

    entries = []
    n_img_rounds = 0
    with open(args.input, encoding="utf-8") as f:
        for line in f:
            r = json.loads(line)
            if args.user and r.get("user_name") != args.user:
                continue
            u = r.get("user_input") or {}
            images = extract_images(u, img_dir, r.get("round_index", 0), desc_map)
            if images:
                n_img_rounds += 1
            entries.append({
                "user_name": r.get("user_name", ""),
                "datetime": r.get("datetime", ""),
                "user_text": u.get("text", ""),
                "images": images,
                "tools": r.get("tool_calls") or [],
                "response": strip_prefix(r.get("response", "")),
            })

    wb = Workbook()
    ws = wb.active
    ws.title = "完整对话记录"
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(1, i)
        c.value = h
        c.font = Font(bold=True)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w

    row = 2
    for e in entries:
        ts = normalize_ts(e["datetime"])
        imgs = e["images"]
        fill = IMG_FILL if imgs else USER_FILL
        note = f"含图片({len(imgs)}张)" if imgs else None

        for cidx, val in [(1, e["user_name"]), (2, "用户"), (3, ts),
                          (4, e["user_text"]), (5, note)]:
            cell = ws.cell(row, cidx)
            cell.value = val
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if imgs:
            ws.cell(row, 6).value = "\n".join(
                f"{i+1}. {im['desc'] or '（见该轮 AI 回复）'}" for i, im in enumerate(imgs))
            ws.cell(row, 7).value = "\n".join(im["name"] for im in imgs)
            pv = preview_dir / f"row{row}.jpg"
            w, h = build_preview(imgs, pv)
            img = ExcelImage(str(pv))
            img.width, img.height = w, h
            ws.add_image(img, f"H{row}")
            ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 18,
                                                h * 0.75 + 8)
        row += 1

        if args.detailed:
            for t in e["tools"]:
                name = t.get("tool_name", "?")
                ok = "✓" if t.get("success") else "✗"
                visible = t.get("visible_text") or t.get("error") or ""
                if len(visible) > 2000:
                    visible = visible[:2000] + "…"
                cells = [(1, e["user_name"]), (2, "工具调用"), (3, ts),
                         (4, visible), (5, f"{name} | {ok}")]
                for cidx, val in cells:
                    cell = ws.cell(row, cidx)
                    cell.value = val
                    cell.fill = TOOL_FILL
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                row += 1

        for cidx, val in [(1, e["user_name"]), (2, "AI回复"), (3, ts), (4, e["response"])]:
            cell = ws.cell(row, cidx)
            cell.value = val
            cell.fill = AI_FILL
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{row - 1}"
    wb.save(args.output)
    print(f"saved={args.output} rows={row - 1} image_rounds={n_img_rounds}")


if __name__ == "__main__":
    sys.exit(main())